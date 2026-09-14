# -*- coding: utf-8 -*-
"""Rebuild US sector workforce observations from archived official tables.
Only writes research/employment/us. No extrapolated values masquerade as observations.
"""
from pathlib import Path
import json,csv,re,hashlib,collections
import openpyxl
ROOT=Path(__file__).resolve().parents[3]; OUT=Path(__file__).resolve().parent; RAW=OUT/'originals'
def read(p):return json.loads((ROOT/p).read_text())
def save(name,obj):(OUT/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n')
def norm(s):return re.sub(r'[^a-z0-9]','',re.sub(r'\\\d+\\','',s.lower()))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
CHECKED='2026-09-14'
inv=[r for r in read('research/employment/node-inventory.json')['records'] if r['country']=='us']
old=read('research/us/subindustry-productivity.json'); annual=read('research/annual-2024-2025/annual-details.json'); macro=read('research/us/macro.json')
oldrows={(r['stablechildId'],r['year']):r for p in old['parents'] for r in p['rows']}
annualrows={(n['id'],b['year']):n for b in annual['branches'] if b['country']=='us' for n in b['nodes']}
macroinds={i['id']:i for i in macro['industries']}
book=openpyxl.load_workbook(RAW/'Section6All.xlsx',data_only=True)
nipa={}; fte={}
for dest,sheet in [(nipa,'T60400D-A'),(fte,'T60500D-A')]:
 s=book[sheet]
 for cells in list(s.rows)[8:105]:
  if cells[0].value and str(cells[0].value).isdigit():
   line=int(cells[0].value); name=re.sub(r'\\\d+\\','',cells[1].value).strip()
   dest[line]={'name':name,'sheet':sheet,'line':line,'years':{int(s.cell(8,col).value):{'value':cells[col-1].value*1000,'cell':cells[col-1].coordinate} for col in range(27,31)}}
concord=openpyxl.load_workbook(ROOT/'research/annual-2024-2025/us/BEA-concordance.xlsx',data_only=True).active
bea={}
for i,row in enumerate(list(concord.values)[5:],6):
 for idx in [1,3,5]:
  if row[idx]:bea.setdefault(norm(str(row[idx])),[]).append({'row':i,'code':str(row[11]),'group':str(row[idx-1]),'level':idx})
xwalk=openpyxl.load_workbook(RAW/'2022_to_2017_NAICS.xlsx',data_only=True).active
new_old=collections.defaultdict(set);old_new=collections.defaultdict(set);crossrows=collections.defaultdict(list)
for i,r in enumerate(list(xwalk.values)[3:],4):
 if isinstance(r[0],(int,float)) and isinstance(r[2],(int,float)):
  new,prior=str(int(r[0])),str(int(r[2]));new_old[new].add(prior);old_new[prior].add(new);crossrows[new].append(i)
q={};qrows={};prefix_cache={}
for y in range(2021,2026):
 records=list(csv.DictReader((RAW/f'qcew-{y}-us.csv').open()))
 q[y]={(r['own_code'],r['industry_code']):r for r in records if r['size_code']=='0'}
 qrows[y]={(r['own_code'],r['industry_code']):i+2 for i,r in enumerate(records) if r['size_code']=='0'}
latest=list(csv.DictReader((RAW/'qcew-2026q1-us.csv').open()))
qlatest={(r['own_code'],r['industry_code']):r for r in latest if r['size_code']=='0'}
latestrow={(r['own_code'],r['industry_code']):i+2 for i,r in enumerate(latest) if r['size_code']=='0'}
# Canonical BEA names differ slightly across workbooks, never match by an ID alone.
alias={
 norm('Tobacco product manufacturing'):'Tobacco manufacturing',
 norm('Wireless telecommunications carriers (except satellites)'):'Wireless telecommunications carriers (except satellite)',
 norm('Real estate (includes owner-occupied housing)'):'Real estate',
 norm('Professional, scientific, and technical services'):'Professional, scientific, and technical services',
}
# BEA summary-to-NIPA employee table correspondence, after checking named rows and notes.
nipa_alias={
 norm('Agriculture, forestry, fishing, and hunting'):4,norm('Construction'):12,norm('Manufacturing'):13,
 norm('Wholesale trade'):35,norm('Retail trade'):38,norm('Information'):52,norm('Finance and insurance'):57,
 norm('Real estate and rental and leasing'):62,norm('Professional, scientific, and technical services'):65,norm('Health care and social assistance'):74,
 norm('Government'):86,norm('Farms'):5,norm('Real estate'):63,
 norm('Publishing industries, except internet (includes software)'):53,
 norm('Data processing, internet publishing, and other information services'):56,
 norm('Wholesale trade, durable goods'):36,norm('Wholesale trade, nondurable goods'):37,
}
for line,r in nipa.items():nipa_alias.setdefault(norm(r['name']),line)
# Federal defense is NOT synonymous with military. No such substitution is made.
government_lines={'us-government-bea-90':87,'us-government-bea-91':88,'us-government-bea-94':91,'us-government-bea-95':92,'us-government-bea-96':93,'us-government-uva-185':94,'us-government-bea-97':96}
SPECIAL_HOUSING={'Housing','Owner-occupied housing','Tenant-occupied housing','Other real estate'}
RELEASE={2021:'2022-08-24',2022:'2023-08-23',2023:'2024-08-21',2024:'2025-09-09',2025:'2026-08-28'}
# Only 2024/2025 publication dates independently located this round; earlier dates remain unknown.
RELEASE.update({2021:None,2022:None,2023:None})
def ev(source,locator):return {'sourceId':source,'locator':locator}
def nipa_emp(line,y,fulltime=False):
 r=(fte if fulltime else nipa)[line];v=r['years'].get(y)
 if not v:return None
 return {'value':v['value'],'unit':'FTE' if fulltime else 'jobs','year':y,'denominatorKind':'fte-employees' if fulltime else 'employee-jobs','periodBasis':'annual-reported','definition':'BEA 年度全职等价雇员量，兼职按工时折算；非自然人数。' if fulltime else 'BEA 年度全职及兼职雇员岗位；一人有多份工作可重复计数，排除业主。','coverage':'BEA 国内行业账户雇员口径；含覆盖调整。全国采用 Domestic industries（第2行），政府包括军人和政府企业；私人医疗不含政府医院。','coverageKey':'us-bea-nipa-domestic-'+str(line),'releaseDate':'2025-09-26','revision':'当前在线工作簿，2026-09-14 下载；表内 Data published September 26, 2025，不能把文件创建日当发布日期。','evidence':[ev('us-employment-bea-nipa-2025-09-26',f"{r['sheet']}!{v['cell']}; line {line} {r['name']}; thousands × 1000; A3 annual 1998—2024; A5 release; row 8 year"),ev('us-employment-nipa-methods', 'Chapter 10 Appendix A, printed 10-25—10-28 / physical 25—28; full/part-time positions, FTE and coverage adjustments')],'periodNote':'表内只明示年度值。手册“annual sum”与月度来源的表述未明确年均处理；保留年度已发布雇员观测，不改写为12个月年均。'}

def node_en(rec):
 k=(rec['nodeId'],rec['year'])
 if k in annualrows:return annualrows[k]['nameEn']
 if k in oldrows:return oldrows[k]['nameEn']
 if rec['nodeId'] in macroinds:return macroinds[rec['nodeId']]['nameEn']
 return 'Domestic industries' if rec['nodeId']=='us-orbit-economy' else rec['name']

def qcew_codes(rec,name,year):
 node=rec['nodeId'];own='5';refs=[]
 if node=='us-orbit-economy':return [('0','10')],refs,[]
 if node=='us-orbit-other':
  selected=set(next(r for r in inv if r['nodeId']=='us-orbit-economy' and r['year']==year)['children'])
  # Remainder is exactly the nine omitted private industries, with classifications checked separately.
  omitted=[i for i in macro['industries'] if i['id'] not in selected]
  codes=[]
  for i in omitted:
   sub=dict(rec,nodeId=i['id']);c,e,g=qcew_codes(sub,i['nameEn'],year)
   if g:return [],e,g
   codes+=c;refs+=e
  if ('5','99') in q[year]:codes.append(('5','99'));refs.append(ev(f'us-employment-qcew-{year}','own_code=5, industry_code=99: unclassified private establishments; included only in unexpanded workforce remainder, not assigned to named industries.'))
  return sorted(set(codes)),refs,[]
 if node=='us-government':return [('1','10'),('2','10'),('3','10')],refs,[]
 if node.startswith('us-government'):
  if name=='Federal':return [('1','10')],refs,[]
  if name=='State and local':return [('2','10'),('3','10')],refs,[]
  if name=='State and local government educational services':return [('2','61'),('3','61')],refs,[]
  return [],refs,['QCEW 按所有权与 NAICS，不能把一般政府/政府企业、国防/非国防及医疗等 BEA 核算拆分等同为相应 NAICS；待 BEA 对应就业细分。']
 if name in SPECIAL_HOUSING:return [],refs,['住房服务与其他房地产按核算活动拆分，531 岗位不能完整分配；自有住房虚拟服务没有可直接配对的雇员分母。']
 if name=='Real estate':raw=['531']
 elif name=='Manufacturing':raw=['31','32','33']
 elif node in macroinds:
  code=macroinds[node]['naicsSectorGroup'];raw=code.split('-') if code not in ['31-33','44-45','48-49'] else {'31-33':['31','32','33'],'44-45':['44','45'],'48-49':['48','49']}[code]
 elif name=='Construction of buildings':raw=['236']
 elif name=='Heavy and civil engineering construction':raw=['237']
 elif name=='Specialty trade contractors':raw=['238']
 elif norm(name)==norm('Merchant wholesalers, durable goods'):raw=['423']
 elif norm(name)==norm('Merchant wholesalers, nondurable goods'):raw=['424']
 elif name in ['Wholesale electronic markets and agents and brokers','Wholesale trade agents and brokers']:raw=['425']
 else:
  rows=bea.get(norm(alias.get(norm(name),name)),[])
  raw=sorted(set(r['code'] for r in rows));refs=[ev('us-employment-bea-concordance', 'NAICS Codes! rows '+','.join(str(r) for r in sorted(set(r['row'] for r in rows)))+'; name '+name+'; column L related 2017 NAICS')]
 if not raw:return [],refs,['已检查 BEA 行业对照，未得到可独立对应的 NAICS 就业分类。']
 if any(not c.isdigit() for c in raw):return [],refs,['BEA 对照列为 '+', '.join(raw)+'，非可直接相配的 NAICS 就业行业；不能按增加值份额推算就业。']
 oldset={c for c in old_new if any(c.startswith(prefix) for prefix in raw)}
 if not oldset:return [],refs,['官方分类对照没有对应完整六位2017分类。']
 if year==2021:target=oldset;universe=set(old_new)
 else:
  target=set().union(*(old_new[c] for c in oldset));partial=[c for c in target if not new_old[c].issubset(oldset)]
  if partial:return [],refs+[ev('us-employment-census-concordance','2022 to 2017 NAICS U.S.; rows '+','.join(str(r) for c in partial for r in crossrows[c]))],['2017→2022 分类拆并不能无权重精确对应；新分类 '+','.join(sorted(partial))+' 混入节点之外的旧分类。需要桥接权重，不能仅凭同名或相似代码相除。']
  universe=set(new_old);refs.append(ev('us-employment-census-concordance','2022 to 2017 NAICS U.S.; all mapped six-digit rows. Selected set checked against every origin industry; mapping.json stores complete lists.'))
 # Use the coarsest available published QCEW rows that exactly cover the target leaf set.
 rem=set(target);codes=[]
 for c in sorted({c for own,c in q[year] if own=='5' and (c.isdigit() and not c.startswith('10') or c in ['31-33','44-45','48-49'])},key=lambda c:(len(c) if '-' not in c else 2,c)):
  if (year,c) not in prefix_cache:
   if '-' in c:lo,hi=c.split('-');prefix_cache[(year,c)]={n for n in universe if int(lo)<=int(n[:2])<=int(hi)}
   else:prefix_cache[(year,c)]={n for n in universe if n.startswith(c)}
  contained=prefix_cache[(year,c)]
  if contained and contained.issubset(rem):codes.append(('5',c));rem-=contained
 if rem:return [],refs,['QCEW 未有完整对应可取数行：'+','.join(sorted(rem))]
 return codes,refs,[]

def make_qcew(codes,year,refs,latest=False):
 d=qlatest if latest else q[year];ridx=latestrow if latest else qrows[year]
 source='us-employment-qcew-2026q1' if latest else f'us-employment-qcew-{year}'
 key='month3_emplvl' if latest else 'annual_avg_emplvl'
 rows=[]
 for own,code in codes:
  r=d.get((own,code))
  if not r or r['disclosure_code'] or not r.get(key):return None
  rows.append({'own_code':own,'industry_code':code,'value':int(r[key]),'csvRow':ridx[(own,code)]})
 if not rows:return None
 return {'value':sum(r['value'] for r in rows),'unit':'jobs','year':2026 if latest else year,'denominatorKind':'covered-employee-jobs','periodBasis':'pay-period' if latest else 'annual-average','definition':'受保雇员岗位（全职、兼职均计；多份工作重复计数），不含自雇业主。','coverage':'QCEW 美国50州及DC；私营行业 own_code=5，政府按联邦/州/地方所有权另取。仅 UI/UCFE 覆盖雇员，排除业主、军人、部分农业/家政/学校及其他非覆盖雇员。','coverageKey':'us-qcew-covered-'+','.join(o+':'+c for o,c in codes),'releaseDate':'2026-08-28' if latest else RELEASE[year],'revision':'2026Q1 初步观测；仅最新进展，不与年度增加值相除。' if latest else '2026-09-14 下载当前在线年度档案；2025年官方年报称 complete and final，历史可因订正更新。','evidence':[ev(source,('2026/1' if latest else str(year)+'/a')+'/area/US000.csv; area_fips=US000; size_code=0; '+('; '.join(f"row {r['csvRow']}, own_code={r['own_code']}, industry_code={r['industry_code']}, {key}={r['value']}" for r in rows))),ev('us-employment-qcew-definition','Employment; Annual Average Employment; Coverage exclusions'),*refs],'calculation':{'kind':'direct-fact' if len(rows)==1 else 'calculation','expression':' + '.join(str(r['value']) for r in rows),'inputs':rows},'periodNote':'2026年3月含12日的工资期岗位数。' if latest else '12个月覆盖雇员岗位数相加÷12；不是全年累计招聘或去重人数。'}

# Census government functional survey provides a separate March employment reference.
# Never sum the national row together with the states contained in this same file.
census_rows=list(csv.DictReader((RAW/'GS00EMP01/GS00EMP01.dat').open(),delimiter='|'))
cnat={(int(r['#YEAR']),r['AGG_DESC']):(i+2,r) for i,r in enumerate(census_rows) if r['GEO_ID']=='0100000US' and r['GOVTYPE']=='001'}
def census_emp(codes,year):
 chosen=[cnat.get((year,c)) for c in codes]
 if not all(chosen) or any(not r['TOT_EMP'].isdigit() or r['TOT_EMP_F'] for _,r in chosen):return None
 inputs=[{'id':c,'label':r['AGG_DESC_LABEL'],'value':int(r['TOT_EMP']),'unit':'jobs','evidence':[ev('us-employment-census-apes-2025',f"GS00EMP01.dat row {line}; YEAR={year}; GEO_ID=0100000US; GOVTYPE=001; AGG_DESC={c}; TOT_EMP={r['TOT_EMP']}; FT_EMP={r['FT_EMP']}, PT_EMP={r['PT_EMP']}")]} for c,(line,r) in zip(codes,chosen)]
 return {'value':sum(x['value'] for x in inputs),'unit':'jobs','year':year,'denominatorKind':'census-government-jobs','periodBasis':'march-pay-period','definition':'Census州及地方政府3月份全职与兼职雇员岗位；按政府功能划分，不是年均或跨雇主去重人数。','coverage':'Census全国州及地方政府调查；按职能分类，未套用BEA一般政府与政府企业的核算调整。医院及卫生功能含公立医院和非医院卫生服务；与BEA核算范围只能作代理配对。','coverageKey':'us-census-apes-march-'+','.join(codes),'releaseDate':'2026-04-16','revision':'2025年首发于2026-04-16；本文件2026-09-14下载，包含历史修订。调查结果随后可能修订。','evidence':[e for i in inputs for e in i['evidence']]+[ev('us-employment-census-apes-summary','2025 Summary, pages 1–2: data relate to March and include full/part-time; release April 16, 2026')],'status':'official' if len(inputs)==1 else 'calculated','typedCalculation':{'methodId':'us-census-march-government-functions','label':'互斥政府功能3月份雇员岗位合计','expression':' + '.join(str(i['value']) for i in inputs),'operation':'identity' if len(inputs)==1 else 'sum','inputs':inputs,'assumptions':[]},'periodNote':'分母是本年3月份调查时点，分子若有值则为本年全年增加值；不是年均人均劳动生产率。'}

records=[]; mappings=[]
for rec in inv:
 node,year=rec['nodeId'],rec['year'];name=node_en(rec);gaps=[];alts=[]
 codes,refs,qgaps=qcew_codes(rec,name,year);covered=make_qcew(codes,year,refs) if codes else None
 mappings.append({'nodeId':node,'year':year,'name':rec['name'],'nameEn':name,'qcewRows':[{'ownership':o,'code':c} for o,c in codes],'evidence':refs,'gaps':qgaps})
 line=2 if node=='us-orbit-economy' else nipa_alias.get(norm(name))
 if node.startswith('us-government') and year>=2024:line=86 if node=='us-government' else government_lines.get(node)
 if name in SPECIAL_HOUSING:line=None
 # Dynamic remainder is not a direct NIPA row. QCEW is disjoint omitted sector aggregate.
 employee=nipa_emp(line,year) if line and year<=2024 else None
 if covered:alts.append(covered)
 if employee:
  f=nipa_emp(line,year,True)
  if f:alts.append(f)
 primary=employee or covered
 if not primary:gaps+=qgaps or ['未取得同年同分类可用就业分母。']
 oldr=oldrows.get((node,year))
 if oldr and oldr['employment']['value'] is not None and oldr['employment'].get('denominatorKind')=='jobs':
  alts.append(dict(oldr['employment'],denominatorKind='employee-and-owner-jobs',unit='jobs',dataOrigin='prior-archived-research',definition=oldr['employment']['definition']))
 # Functional hospital/health employment is useful despite the March/accounting boundary.
 c_codes=['EP0005'] if name=='State and local' else ['EP1085'] if name=='State and local government educational services' else ['EP0605','EP0645'] if name=='State and local government hospitals and health services' else None
 if c_codes:
  cemp=census_emp(c_codes,year)
  if cemp:
   alts.append(cemp)
   if not primary:
    primary=cemp;gaps.append('采用Census医院及卫生功能3月岗位作代理；其时点与BEA核算调整不完全一致，不能与年均NIPA或QCEW人数混排。')
 if node.startswith('us-construction-uva-'):
  gaps.append('已检查2022建筑经济普查BASIC/KOB/LOCCONS/VALCON表目录和调查说明：就业及工时按雇主行业，工程类型公布业务额；未取得工程类型×就业/工时分配矩阵。Construction Spending排除维修维护，不能据此覆盖本图8类活动。')
 special=name in SPECIAL_HOUSING or name=='Customs duties'
 if name=='Customs duties':gaps=['关税是核算调整项，不是有独立就业的经营行业；就业和人均增加值不适用，不能填0后相除。']
 if primary and node=='us-orbit-other' and year==2025:gaps.append('就业余项另含QCEW尚未分类的私营岗位（NAICS 99）；其行业归属未知，未分配给具名行业。该就业覆盖不能解释为恰好9个BEA行业的完整同口径人数。')
 if primary and node.startswith('us-agriculture'):gaps.append('农业业主和部分农业雇员未包含在雇员岗位中；该比率不能代表全口径农业人均产出。')
 if primary and (node=='us-real-estate' or name=='Real estate'):gaps.append('房地产增加值含自有住房估算租金；比率同时反映住房资产服务，不能归因为房地产雇员劳动效率。')
 if primary and node.startswith('us-government') and primary['denominatorKind']=='covered-employee-jobs':gaps.append('政府 QCEW 分母仅覆盖文职雇员，缺军人及其他 BEA 覆盖调整；不要解释为全体政府从业人员效率。')
 va=rec['valueAddedBase'];ratio=None
 if primary and va is not None and primary['value']>0 and not special:
  ratio={'value':va/primary['value'],'unit':'USD/job/year' if primary['unit']=='jobs' else 'USD/FTE/year','status':'proxy','formula':f"{va} USD / {primary['value']} {primary['unit']}",'numeratorPeriod':year,'denominatorPeriod':year,'label':'每雇员岗位对应增加值（估算）' if employee else '每受保雇员岗位对应增加值（代理比率）','assumptions':['分子为本行业全部现价增加值；分母是雇员岗位，未含业主就业。数值描述行业规模与雇员岗位的比值，不是全口径人均劳动生产率。','年度与分类可对应；不把劳动、资本及生产税等共同形成的增加值全部归因为劳动。']+([employee['periodNote']] if employee else [])+gaps}
 elif va is None:gaps.append('所选年份的该细分现价增加值未发布/未取得；即便有就业也不借用另一年产出计算。')
 for alt in alts:
  if va is not None and alt['value']>0 and not special:
   alt['ratio']={'value':va/alt['value'],'unit':'USD/FTE/year' if alt['unit']=='FTE' else 'USD/job/year','status':'proxy' if alt['denominatorKind']!='employee-and-owner-jobs' else 'calculation-with-scope-limitations','formula':f"{va} USD / {alt['value']} {alt['unit']}"}
 latestemp=make_qcew(codes,2025,refs,True) if year==2025 and codes else None
 records.append({'id':f"us-employment-{node}-{year}",'country':'us','nodeId':node,'basis':rec['basis'],'year':year,'name':rec['name'],'nameEn':name,'parentId':rec['parentId'],'kind':rec['kind'],'status':'not-applicable' if name in ['Owner-occupied housing','Customs duties'] else 'available' if primary else 'gap','employment':primary,'valueAdded':{'value':va,'unit':'USD','year':year,'priceBasis':'current','evidence':rec['valueEvidence']},'ratio':ratio,'alternatives':alts,'latestEmployment':latestemp,'gaps':gaps,'evidence':refs})
# Root/remainder need consistent denominator options for comparison; NIPA remainder derives its same-year omitted rows.
for rec in records:
 if rec['nodeId']=='us-orbit-other' and rec['year']<=2024:
  selected=set(next(r for r in inv if r['nodeId']=='us-orbit-economy' and r['year']==rec['year'])['children'])
  other=[i for i in macro['industries'] if i['id'] not in selected]
  lines=[nipa_alias.get(norm(i['nameEn'])) for i in other]
  if all(lines):
   inputs=[nipa_emp(line,rec['year']) for line in lines];emp=dict(inputs[0]);emp['value']=sum(r['value'] for r in inputs);emp['definition']='未展示的9个互斥私人行业年度雇员岗位合计；非去重人数。';emp['coverageKey']='us-bea-nipa-other';emp['evidence']=[e for r in inputs for e in r['evidence']];emp['calculation']={'kind':'calculation','expression':' + '.join(str(r['value']) for r in inputs),'lines':lines};rec['employment']=emp;rec['status']='available';va=rec['valueAdded']['value'];rec['ratio']={'value':va/emp['value'],'unit':'USD/job/year','status':'proxy','formula':f"{va} USD / {emp['value']} jobs",'numeratorPeriod':rec['year'],'denominatorPeriod':rec['year'],'label':'每雇员岗位对应增加值（估算）','assumptions':['未列示的互斥私人行业增加值合计÷相同9行业雇员岗位合计；未含业主，非全口径自然人人均。',emp['periodNote']]}
# Explicit model for classification breaks: same-year parent employee pool, old-classification QCEW shares.
# This estimates current-year EMPLOYEE jobs, not all jobs or unique persons; source-year weights stay visible.
index={(r['nodeId'],r['year']):r for r in records};inventory={(r['nodeId'],r['year']):r for r in inv}
for repeat in range(3):
 for r in records:
  if r['employment'] or r['status']=='not-applicable' or r['kind']=='remainder':continue
  if r['nodeId'].startswith(('us-government','us-construction','us-real-estate')):continue
  parent=index.get((r['parentId'],r['year']))
  if not parent or not parent['employment']:continue
  if not any('2017→2022' in g for g in r['gaps']):continue
  target_rec=dict(inventory[(r['nodeId'],r['year'])],year=2021)
  pool_rec=dict(inventory[(r['parentId'],r['year'])],year=2021)
  tc,te,tg=qcew_codes(target_rec,r['nameEn'],2021)
  pc,pe,pg=qcew_codes(pool_rec,parent['nameEn'],2021)
  tm=make_qcew(tc,2021,te) if tc else None;pm=make_qcew(pc,2021,pe) if pc else None
  if not tm or not pm or not pm['value'] or tm['value']>pm['value']:continue
  pool=parent['employment'];v=pool['value']*tm['value']/pm['value']
  assumptions=['假设本节点占父节点的雇员岗位比重仍等于2021年QCEW旧分类比重；新行业分类本身不提供拆分权重。','将2021年受保雇员结构用于所注明同年父项岗位池；行业增减、业主及非覆盖调整的结构变化可能令估算偏离实际。','估算不与官方实测兄弟项强制配平，不能把二者相加宣称完整就业普查。']
  emp={'value':v,'unit':'jobs','year':r['year'],'denominatorKind':pool['denominatorKind'],'periodBasis':pool['periodBasis'],'status':'estimated','definition':'按2021年旧分类就业份额分配同年父行业雇员岗位（模型估算）；不是官方行业人数。','coverage':pool['coverage'],'coverageKey':pool['coverageKey']+'-model2021','releaseDate':None,'revision':'研究估算，2026-09-14；同年岗位池×2021旧分类就业份额。','evidence':pool['evidence']+tm['evidence']+pm['evidence'],'typedCalculation':{'methodId':'us-prior-qcew-2021-share','label':'同年父行业岗位 × 2021年细分就业份额','expression':f"{pool['value']} * {tm['value']} / {pm['value']}",'operation':'share','inputs':[{'id':'current-parent','label':str(r['year'])+'年 '+parent['name']+'岗位池','value':pool['value'],'unit':'jobs','evidence':pool['evidence']},{'id':'child-2021','label':'2021年 '+r['name']+'受保雇员岗位','value':tm['value'],'unit':'jobs','evidence':tm['evidence']},{'id':'parent-2021','label':'2021年 '+parent['name']+'受保雇员岗位','value':pm['value'],'unit':'jobs','evidence':pm['evidence']}],'assumptions':assumptions},'sensitivity':{'lower':v*0.8,'upper':v*1.2,'label':'旧分类份额相对偏离±20%的情景范围（非置信区间）','assumptions':['仅假设所用旧份额相对点估算上下变动20%，岗位池不变；不是统计误差区间，不能保证覆盖真实值。']}}
  emp['periodNote']=pool.get('periodNote','本年年度雇员观测')
  r['employment']=emp;r['status']='available';r['gaps'].append('已补显式模型；旧分类分拆的官方就业观测仍缺。')
  va=r['valueAdded']['value']
  if va is not None:r['ratio']={'value':va/v,'unit':'USD/job/year','status':'proxy','formula':f"{va} USD / ({pool['value']} * {tm['value']} / {pm['value']}) jobs",'numeratorPeriod':r['year'],'denominatorPeriod':r['year'],'label':'每估算雇员岗位对应增加值','assumptions':assumptions}
exec((OUT/'add_models.py').read_text(),globals())
# A wholly unallocated monetary parent has exactly the same workforce scope as its parent.
for r in records:
 if r['employment'] or r['kind']!='remainder' or r['nodeId']=='us-orbit-other':continue
 parent=index.get((r['parentId'],r['year']))
 if not parent or not parent['employment'] or r['valueAdded']['value']!=parent['valueAdded']['value']:continue
 emp=dict(parent['employment']);emp['status']='estimated' if emp.get('status')=='estimated' else 'calculated';emp.pop('calculation',None);emp['typedCalculation']={'methodId':'us-whole-parent-remainder','label':'全部父项金额未分列，就业范围等于父项','expression':str(emp['value']),'operation':'identity','inputs':[{'id':'parent-employment','label':parent['name']+'同年就业','value':emp['value'],'unit':'jobs','evidence':emp['evidence']}],'assumptions':['仅当余项增加值等于父项完整增加值时采用同一就业规模。该余项不是新增就业，不能与父项重复加总。']};r['employment']=emp;r['status']='available';r['gaps']=['全部父项金额尚未分列；就业沿用同范围父项，不分配给各个未知子项。'];va=r['valueAdded']['value'];r['ratio']={'value':va/emp['value'],'unit':'USD/job/year','status':'proxy','formula':f"{va} / {emp['value']}",'numeratorPeriod':r['year'],'denominatorPeriod':r['year'],'label':'未分列父项每雇员岗位增加值','assumptions':emp['typedCalculation']['assumptions']}
sources=[]
def src(sid,title,url,file,published=None,publisher='U.S. Bureau of Labor Statistics',limitations=[]):
 p=RAW/file;sources.append({'id':sid,'country':'us','title':title,'publisher':publisher,'url':url,'published':published,'retrieved':CHECKED,'kind':'official-statistics','archive':str(p.relative_to(ROOT)),'sha256':sha(p),'limitations':limitations})
for y in range(2021,2026):src(f'us-employment-qcew-{y}',f'QCEW {y} annual average employment — U.S. all industries and ownerships',f'https://data.bls.gov/cew/data/api/{y}/a/area/US000.csv',f'qcew-{y}-us.csv',RELEASE[y],limitations=['仅 UI/UCFE 覆盖雇员岗位；全国合计不含波多黎各和美属维尔京群岛。','历史发布日期未核实的保留空值；不能将下载日当发布日期。'] if not RELEASE[y] else ['仅 UI/UCFE 覆盖雇员岗位；全国合计不含波多黎各和美属维尔京群岛。'])
src('us-employment-qcew-2026q1','QCEW first quarter 2026 — U.S. employment','https://data.bls.gov/cew/data/api/2026/1/area/US000.csv','qcew-2026q1-us.csv','2026-08-28',limitations=['2026年3月工资期岗位是最新进展，不能与2025全年产出相除。'])
src('us-employment-bea-io2017','BEA Supply-Use 2017 detailed benchmark compensation by activity','https://apps.bea.gov/industry/release/zip/SUPPLY-USE.zip','SUPPLY-USE.zip',None,'U.S. Bureau of Economic Analysis',['2017年基准薪酬结构，不是当前就业；工程/住房分配假设同类年均每岗位薪酬相同。工作簿未标发布日期，保持空值。'])
src('us-employment-census-apes-2025','Annual Survey of Public Employment and Payroll, national functional data 1993—2025','https://www2.census.gov/programs-surveys/apes/data/GS00EMP01.zip','GS00EMP01.zip','2026-04-16','U.S. Census Bureau',['按政府功能与3月调查时点，不是BEA年度国内雇员账户；本表同时含全国和州记录，提取仅选全国。'])
src('us-employment-census-apes-summary','2025 Annual Survey of Public Employment and Payroll Summary Report','https://www.census.gov/content/dam/Census/library/publications/2025/econ/ASPEP%20Summary%20Report%202025.pdf','aspep2025-summary.pdf','2026-04-16','U.S. Census Bureau',['方法与发布日期参考；具体人数取完整批量表，未用摘要四舍五入值。'])
src('us-employment-bea-nipa-2025-09-26','NIPA 6.4D employees and 6.5D FTE, annual 1998—2024','https://apps.bea.gov/national/Release/XLS/Survey/Section6All_xls.xlsx','Section6All.xlsx','2025-09-26','U.S. Bureau of Economic Analysis',['不含业主；年度统计时点保留表内年度观测，不将手册未说明的年均处理写成事实。'])
src('us-employment-nipa-methods','NIPA Handbook Chapter 10, Appendix A employment definitions','https://www.bea.gov/resources/methodologies/nipa-handbook/pdf/chapter-10.pdf','chapter-10.pdf',None,'U.S. Bureau of Economic Analysis',['本轮读取附录A及覆盖调整，非整本NIPA手册。'])
src('us-employment-census-concordance','2022 NAICS matched to 2017 NAICS — full concordance','https://www.census.gov/naics/concordances/2022_to_2017_NAICS.xlsx','2022_to_2017_NAICS.xlsx',None,'U.S. Census Bureau',['分类表无就业权重；一对多、多对一遇到跨行业边界时不强行分配。'])
src('us-employment-qcew-definition','Employment and Wages Annual Averages 2025 — coverage and definitions','https://www.bls.gov/cew/publications/employment-and-wages-annual-averages/2025/home.htm','qcew-definition-web.json','2026-08-28',limitations=['网页归档为Web工具可读摘录，非原始HTML。'])
src('us-employment-qcew-notices','QCEW notices — annual release dates','https://www.bls.gov/cew/notices/','qcew-notices-web.json',None,limitations=['网页归档为Web工具可读摘录，非原始HTML。'])
src('us-employment-qcew-geography','QCEW U.S. total excludes Puerto Rico and Virgin Islands — footnote','https://www.bls.gov/regions/northeast/news-release/2026/countyemploymentandwages_puertorico_20260612.htm','qcew-geography-web.json','2026-06-12',limitations=['网页归档为Web工具可读摘录，非原始HTML。'])
src('us-employment-qcew-naics','QCEW industry classification history — 2017 through 2021, 2022 onward','https://www.bls.gov/cew/classifications/industry/industry-titles.htm','qcew-naics-web.json',None,limitations=['网页归档为Web工具可读摘录，非原始HTML。'])
src('us-employment-bea-discontinued','BEA discontinued and delayed statistics','https://www.bea.gov/data/discontinued-or-delayed-statistics','bea-discontinued.htm',None,'U.S. Bureau of Economic Analysis',['SAEMP及NIPA自雇和生产参与者表停发；不能将停发之后模型值标为官方观测。'])
conpath=ROOT/'research/annual-2024-2025/us/BEA-concordance.xlsx'
sources.append({'id':'us-employment-bea-concordance','title':'BEA Industry and Commodity Codes and 2017 NAICS Concordance','country':'us','publisher':'U.S. Bureau of Economic Analysis','url':'https://www.bea.gov/sites/default/files/2023-10/BEA-Industry-and-Commodity-Codes-and-NAICS-Concordance.xlsx','published':None,'retrieved':CHECKED,'kind':'official-statistics','archive':str(conpath.relative_to(ROOT)),'sha256':sha(conpath),'limitations':['住房/工程类型/关税及政府活动不是独立NAICS雇主行业。']})
# Reused numerical VA and historical denominator sources retain original objects and immutable provenance.
refs={e['sourceId'] for r in records for e in r['valueAdded']['evidence']}
for r in records:
 for a in r['alternatives']:refs.update(e['sourceId'] for e in a.get('evidence',[]))
known={s['id'] for s in sources}
for s in old['sources']+annual['sources']:
 if s['id'] in refs and s['id'] not in known:sources.append(s);known.add(s['id'])
# Main macro sources may have another schema; existing integration has source IDs.
coverage=[{'country':'us','nodeId':r['nodeId'],'year':r['year'],'basis':r['basis'],'name':r['name'],'status':r['status'],'employmentAvailable':r['employment'] is not None,'ratioAvailable':r['ratio'] is not None,'reason':r['gaps'],'denominator':r['employment']['denominatorKind'] if r['employment'] else None} for r in records]
summary={'checkedAt':CHECKED,'nodeYears':len(records),'nodeIds':len({r['nodeId'] for r in records}),'employmentAvailable':sum(r['employment'] is not None for r in records),'ratios':sum(r['ratio'] is not None for r in records),'byYear':{str(y):{'total':sum(r['year']==y for r in records),'employment':sum(r['year']==y and r['employment'] is not None for r in records),'ratio':sum(r['year']==y and r['ratio'] is not None for r in records)} for y in range(2021,2026)}}
save('sources.json',sources);save('observations.json',records);save('coverage.json',coverage);save('mapping.json',mappings);save('summary.json',summary)
assert len({(r['nodeId'],r['year'],r['basis']) for r in records})==len(inv)
assert all(not r['ratio'] or (r['valueAdded']['value'] is not None and r['employment']['year']==r['year']) for r in records)
assert all(not r['employment'] or r['employment']['value']>0 for r in records)
print(json.dumps(summary,ensure_ascii=False,indent=2))

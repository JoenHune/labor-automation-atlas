"""Reproduce the US 2024–2025 VA candidate from the saved official workbooks.

Requires openpyxl. Writes only in this directory. No estimates or employment ratios.
"""
import hashlib
import json
from collections import defaultdict
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
CHECKED = "2026-09-14"
FILES = {
    "main": ("ValueAdded-2026-09-14.xlsx", "TVA105-A", "2026-06-25", "https://apps.bea.gov/industry/Release/XLS/GDPxInd/ValueAdded.xlsx"),
    "detail": ("UnderlyingValueAdded-2026-09-14.xlsx", "UVA205-A", "2025-09-25", "https://apps.bea.gov/industry/Release/XLS/UGdpxInd/ValueAdded.xlsx"),
}

TRANSLATIONS = {
 5:"种植业", 6:"畜牧业与水产养殖", 16:"教育、医院与医疗建筑", 17:"建筑维修与维护",
 18:"办公与商业建筑", 19:"其他住宅建设", 20:"其他非住宅建筑", 21:"电力与通信设施建设",
 22:"独户住宅建设", 23:"交通设施、公路与街道建设", 25:"耐用品制造", 56:"非耐用品制造",
 29:"钢铁冶炼与购入钢材加工", 30:"有色金属生产加工与铸造", 33:"农业机械制造",
 34:"建筑机械制造", 35:"采矿与油气田机械制造", 36:"其他机械制造",
 38:"计算机及外围设备制造", 39:"通信设备制造", 40:"半导体与其他电子元件制造",
 41:"导航、测量、医用电子及控制仪器制造", 42:"其他计算机与电子产品制造",
 45:"轿车制造", 46:"轻型卡车与多用途车辆制造", 47:"重型卡车制造", 48:"汽车车身、挂车与零部件制造",
 50:"航空航天产品与零部件制造", 51:"其余交通运输设备制造", 54:"医疗设备与用品制造",
 55:"其他杂项制造", 58:"食品制造", 59:"饮料制造", 60:"烟草制品制造",
 67:"基础化学品制造", 68:"树脂、橡胶与人造纤维制造", 69:"药品与医药制造", 70:"其他化学品制造",
 73:"汽车、汽车零部件及用品批发", 74:"专业与商业设备及用品批发",
 75:"家电、电气与电子产品批发", 76:"机械设备及用品批发", 77:"其他耐用品批发",
 78:"药品及相关用品批发", 79:"食品及相关产品批发", 80:"石油及石油产品批发",
 81:"其他非耐用品批发", 82:"批发电子市场、代理与经纪", 83:"海关关税（BEA 批发业核算项）",
 89:"建材、园艺设备与用品零售", 90:"健康与个人护理用品零售", 91:"加油站",
 92:"服装与配饰零售", 93:"无店铺零售", 94:"其余零售",
 108:"报纸、期刊、图书与目录出版", 109:"软件出版", 112:"广播电视（互联网除外）",
 113:"有线电信运营", 114:"无线电信运营（卫星除外）", 115:"其他电信（含卫星）",
 117:"数据处理、托管与相关服务", 118:"其他信息服务", 124:"直接人寿保险",
 125:"其他保险承保", 126:"保险代理、经纪与相关服务", 130:"住房服务", 131:"业主自住住房服务",
 132:"出租住房服务", 133:"其他房地产业", 140:"会计、报税、记账与工资服务",
 141:"建筑、工程及相关专业服务", 142:"管理、科学与技术咨询", 143:"科学研发服务",
 144:"广告、公关及相关服务", 145:"专业设计及其他专业科技服务", 157:"医师诊所",
 158:"牙医诊所", 159:"其他健康从业者诊所", 160:"门诊护理中心", 161:"其他门诊医疗服务",
 179:"联邦一般政府", 180:"国防", 181:"非国防", 182:"联邦政府企业",
 184:"州与地方一般政府", 185:"州与地方公立教育服务", 186:"州与地方公立医院及医疗服务",
 187:"州与地方政府其他服务", 188:"州与地方政府企业",
}

def canonical(name):
    return {"Federal general government": "General government", "State and local general government": "General government"}.get(name, name)

def read_table(kind, last_line):
    file, sheet, release, url = FILES[kind]
    ws = openpyxl.load_workbook(HERE / file, data_only=True)[sheet]
    assert ws["A2"].value == "[Millions of dollars]"
    assert ws["A5"].value == {"main": "Data published June 25, 2026", "detail": "Data published September 25, 2025"}[kind]
    years = {int(c.value): c.column for c in ws[8] if str(c.value).isdigit()}
    stack, rows = [], {}
    for row in ws:
        if not str(row[0].value).isdigit() or not (2 <= int(row[0].value) <= last_line):
            continue
        line, name = int(row[0].value), row[1].value
        indent = len(name) - len(name.lstrip())
        label = canonical(name.strip())
        while stack and stack[-1][0] >= indent:
            stack.pop()
        parent = stack[-1][2] if stack else None
        key = "/".join([v[1] for v in stack] + [label])
        rows[key] = {"line": line, "nameEn": name.strip(), "indent": indent, "key": key, "parentKey": parent, "row": row[0].row,
                     "values": {year: ws.cell(row[0].row, col).value for year, col in years.items()},
                     "cells": {year: ws.cell(row[0].row, col).coordinate for year, col in years.items()}}
        stack.append((indent, label, key))
    return rows, ws, years

main, main_ws, main_years = read_table("main", 97)
detail, detail_ws, detail_years = read_table("detail", 188)
assert max(detail_years) == 2024 and max(main_years) == 2025
overlap = []
for key, row in main.items():
    other = detail[key]
    overlap.append({"nameEn": row["nameEn"], "path": key, "year": 2024,
                    "mainCell": row["cells"][2024], "detailCell": other["cells"][2024],
                    "mainValue": row["values"][2024], "detailValue": other["values"][2024],
                    "equal": row["values"][2024] == other["values"][2024]})
assert all(r["equal"] for r in overlap)
assert main_ws["AE9"].value == detail_ws["AE9"].value

macro = json.loads((REPO / "research/us/macro.json").read_text())
old = json.loads((REPO / "research/us/subindustry-productivity.json").read_text())
selected_ids = {p["parentIndustryId"] for p in old["parents"]}
top = {i["beaLine"]: i for i in macro["industries"] if i["id"] in selected_ids}
roots = {k: top[r["line"]] for k,r in main.items() if r["line"] in top}
old_names = {r["stablechildId"]: r["name"] for p in old["parents"] for r in p["rows"]}
chosen = {k:r for k,r in detail.items() if any(k == root or k.startswith(root + "/") for root in roots)}
ids, top_ids = {}, {}
for key, row in chosen.items():
    root_key = next(k for k in roots if key == k or key.startswith(k + "/"))
    root = roots[root_key]
    top_ids[key] = root["id"]
    ids[key] = root["id"] if key == root_key else (f'{root["id"]}-bea-{main[key]["line"]}' if key in main else f'{root["id"]}-uva-{row["line"]}')

conc = openpyxl.load_workbook(HERE / "BEA-concordance.xlsx", data_only=True).worksheets[0]
concordance = defaultdict(list)
for row in list(conc)[5:]:
    for code_col, name_col, level in [(0,1,"sector"),(2,3,"summary"),(4,5,"underlying-summary")]:
        if row[name_col].value:
            concordance[str(row[name_col].value).strip().lower()].append({"code":str(row[code_col].value), "level":level,
               "row":row[0].row,"codeCell":row[code_col].coordinate,"descriptionCell":row[name_col].coordinate,
               "naics":str(row[11].value), "notes":row[10].value})

sources = []
for kind, (file,sheet,release,url) in FILES.items():
    sources.append({"id":f"us-bea-{kind}-value-added-{release}","title":"Value Added by Industry" if kind=="main" else "U.Value Added by Industry — Underlying Detail",
       "country":"us","publisher":"U.S. Bureau of Economic Analysis","kind":"official-statistics","url":url,"file":file,"sheet":sheet,
       "releaseDate":release,"checkedAt":CHECKED,"latestDataYear":2025 if kind=="main" else 2024,"unit":"USD million","priceBasis":"current",
       "sha256":hashlib.sha256((HERE/file).read_bytes()).hexdigest(),
       "revision":"latest available official table vintage at check; subject to revision",
       "qualityNote":None if kind=="main" else "BEA 明示细分估计质量显著低于其所属上层汇总，见 UVA205-A!A205。"})

nodes, observations = [], []
for key, row in chosen.items():
    node_id, root_id = ids[key], top_ids[key]
    match = concordance[row["nameEn"].lower()]
    seen = set()
    match = [m for m in match if not ((m["code"],m["level"],m["row"]) in seen or seen.add((m["code"],m["level"],m["row"]))) ]
    node = {"id":node_id,"country":"us","parentIndustryId":root_id,"parentId":ids.get(row["parentKey"]),
        "name": roots[key]["name"] if key in roots else old_names.get(node_id, TRANSLATIONS.get(row["line"],row["nameEn"])),
        "nameEn":row["nameEn"],"code": f'UVA205:{row["line"]}', "beaUnderlyingLine":row["line"],
        "beaMainLine": main[key]["line"] if key in main else None,
        "classification":"BEA GDP by Industry / 2017 NAICS related classification",
        "hierarchyLocator":f'UVA205-A!B{row["row"]}; indentation {row["indent"]}; parent row in same table',
        "concordance":match,
        "officialCodes":sorted(set(m["code"] for m in match)),
        "categoryKind":"tax-accounting-component" if row["line"]==83 else "construction-type" if 16<=row["line"]<=23 else "industry",
        "notes":[]}
    if row["line"]==83: node["notes"].append("海关关税是 BEA 列在批发业下的核算项，收入归政府；不是批发企业收入或独立经营行业，亦不是政府行业增加值。")
    if 16<=row["line"]<=23: node["notes"].append("按 BEA 工程类型分类；不可与 NAICS 236/237/238 并列或按名称一一映射。")
    if row["line"] in (130,131,132): node["notes"].append("住房服务包括业主自住估算租金；不代表房地产雇员独立产出。")
    # This display hint preserves the existing 19 manufacturing siblings while retaining official aggregate nodes.
    node["displayParentId"] = root_id if row["parentKey"] in chosen and chosen[row["parentKey"]]["line"] in (25,56) else node["parentId"]
    node["optionalGroupingNode"] = row["line"] in (25,56)
    nodes.append(node)
    for year in (2024,2025):
        kind = "main" if key in main else "detail"
        source_row = main[key] if key in main else row
        file,sheet,release,url=FILES[kind]
        exists = year in source_row["values"]
        value = source_row["values"].get(year)
        assert value is None or isinstance(value,(int,float))
        loc = {"file":file,"sheet":sheet,"beaLine":source_row["line"],"cell":source_row["cells"].get(year),
               "periodHeaderCell":f'{openpyxl.utils.get_column_letter((main_years if kind=="main" else detail_years)[year])}8' if exists else None,
               "titleCell":"A1","unitCell":"A2","releaseDateCell":"A5","availableYearsCell":"A3"}
        observations.append({"nodeId":node_id,"country":"us","parentIndustryId":root_id,"parentId":node["parentId"],"code":node["code"],
            "year":year,"value":value,"unit":"USD million","currency":"USD","measure":"value-added","priceBasis":"current",
            "frequency":"annual","annualRate":False,"status":"direct-fact" if exists else "not-yet-published",
            "releaseDate":release if exists else None,"revision":"latest available official vintage; subject to revision" if exists else "no observation published in the checked table",
            "checkedAt":CHECKED,"coverage":"BEA 美国国内行业账户；私人行业与政府（含政府企业）按官方分类分列，建筑按工程类型；批发含独列关税核算项。",
            "sourceId":f"us-bea-{kind}-value-added-{release}","sourceUrl":url,"sourceLocator":loc,
            "sourceExcerpt":f'{source_row["nameEn"]}; {year}; {value}; [Millions of dollars]' if exists else "Annual data from 1997 to 2024",
            "evidenceLocator":f'{sheet}!{loc["cell"]}; year {loc["periodHeaderCell"]}; unit A2; release A5' if exists else f'{sheet}!A3 + D8:AE8 (latest year 2024; 2025 not present)',
            "qualityNote":None if kind=="main" else "BEA：细分估计质量显著低于上层汇总（UVA205-A!A205）。",
            "gap":None if exists else "当前官方 Underlying Detail 末年为 2024，未发布 2025；不沿用旧值、不按旧份额外推。"})

by_id = {n["id"]:n for n in nodes}
for node in nodes:
    node["parentCode"] = by_id[node["parentId"]]["code"] if node["parentId"] else None
for observation in observations:
    node = by_id[observation["nodeId"]]
    observation["officialCodes"] = node["officialCodes"]
    observation["parentCode"] = node["parentCode"]
values = {(o["nodeId"],o["year"]):o["value"] for o in observations}
checks=[]
for parent in nodes:
    children=[n for n in nodes if n["parentId"]==parent["id"]]
    if not children:continue
    for year in (2024,2025):
        child_values=[values[(n["id"],year)] for n in children]
        if None in child_values:continue
        total=values[(parent["id"],year)]
        delta=sum(child_values)-total
        # Integers in million-dollar table can differ by the accumulated half-unit rounding bounds.
        bound=(len(children)+1)/2
        assert abs(delta)<=bound,(parent["name"],year,delta)
        checks.append({"parentId":parent["id"],"year":year,"parentValue":total,"childIds":[n["id"] for n in children],
           "childSum":sum(child_values),"delta":delta,"unit":"USD million","roundingBound":bound,"withinRoundingBound":True})

result={"version":"2026-09-14-us-annual-va-01","country":"us","checkedAt":CHECKED,
    "scope":"11 个已入选父行业；2024/2025 现价年度增加值；官方可识别层级（未拆分汇总原样保留）",
    "sources":sources,"nodes":nodes,"observations":observations,
    "availability":{"2024":{"nodesWithValue":sum(o["year"]==2024 and o["value"] is not None for o in observations)},
                    "2025":{"nodesWithValue":sum(o["year"]==2025 and o["value"] is not None for o in observations),"unpublishedDetailNodes":sum(o["year"]==2025 and o["value"] is None for o in observations)}},
    "integrationNotes":["保留现有 us-*-bea-* 稳定 ID；新增细表节点使用 us-*-uva-*，不复用已经表示不同分类的 NAICS ID。",
       "2024 重叠父值与最新版主表逐值完全相等；细表叶值仍注明其自身 2025-09-25 发布版本，不冒写成 2026-06-25。",
       "parentId 保留官方层级；optionalGroupingNode 的耐用品/非耐用品节点可不显示，此时直属子项使用 displayParentId，不能将分组总计与其子项同榜重复相加。",
       "批发 11 项中的海关关税必须独列并解释核算性质；不可把它填入某个 NAICS 423/424/425 行业。",
       "建筑 8 项是 BEA 工程类型划分，与旧 NAICS 236/237/238 行业并不一一对应，应替换视图的分解方案而不是混排。",
       "2025 空值保留缺口，不能静默沿用 2024 或 2023；后续发布日以官方当前安排 2026-09-30 为准，发布后须重新下载核查。"]}
(HERE/"us-annual-va-candidate.json").write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n")
validation={"checkedAt":CHECKED,"overlappingRowsCompared":len(overlap),"allOverlapEqual":True,"overlap":overlap,
   "gdp2024Match":{"mainCell":"TVA105-A!AE9","detailCell":"UVA205-A!AE9","value":main_ws["AE9"].value},
   "parentChildChecks":checks,"allParentChildChecksWithinRounding":True,"availability":result["availability"],
   "note":"空值子项没有伪装成合计通过；仅比对数值完整且非重叠的官方直接父子组。"}
(HERE/"validation.json").write_text(json.dumps(validation,ensure_ascii=False,indent=2)+"\n")
print(json.dumps({"nodes":len(nodes),"availability":result["availability"],"overlap":len(overlap),"parentChildChecks":len(checks)},ensure_ascii=False))

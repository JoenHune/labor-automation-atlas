"""Audit BEA's live-response archives for 2025 detail availability.

Writes only in this follow-up directory. No estimates and no site mutation.
"""
from pathlib import Path
from html import unescape
import hashlib
import json
import re
import openpyxl

HERE = Path(__file__).resolve().parent
CHECKED = "2026-09-14"

def prompts(name):
    result = json.loads((HERE/f"{name}-response.json").read_text())
    return {p["Name"]:p for p in result["Steps"][0]["Prompts"]}

def interact(name):
    values = prompts(name)
    table = json.loads(json.loads(values["TheTable"]["PromtData"])["Table"])
    headers = [c["CV"] for c in table["Data_Rows"][0]]
    years = [int(c) for c in headers if str(c).isdigit()]
    first_choices = json.loads(values["First_Year"]["PromtData"])["Table"]
    last_choices = json.loads(values["Last_Year"]["PromtData"])["Table"]
    assert 2025 not in years
    assert max(int(c["PrdNbr"]) for c in first_choices) == 2024
    assert max(int(c["PrdNbr"]) for c in last_choices) == 2024
    return {"title":table["Title"],"unit":table["Sub_Title"],"revisionText":table["Description"],
            "displayedYears":years,"maximumSelectableYear":2024,"has2025Column":False,
            "rawResponse":f"{name}-response.json","requestFile":f"{name}-request.json",
            "locator":"Steps[0].Prompts: TheTable.Data_Rows[0]; First_Year and Last_Year.PromtData.Table"}

direct = interact("itable-uva-current")
components = interact("itable-components-current")
forced = prompts("itable-2025")
try:
    json.loads(json.loads(forced["TheTable"]["PromtData"])["Table"])
    raise AssertionError("Forced out-of-range response changed; inspect the new structure")
except json.JSONDecodeError:
    malformed = True
assert max(r["PrdNbr"] for r in json.loads(forced["Last_Year"]["PromtData"])["Table"]) == 2024

archive = openpyxl.load_workbook(HERE/"ComponentsOfVA.xlsx",data_only=True)
archive_headers=[]
for sheet in archive:
    if sheet.title == "Contents": continue
    years=[int(c.value) for c in sheet[8] if str(c.value).isdigit()]
    assert max(years)==2023 and 2025 not in years
    assert sheet["A3"].value=="Annual data from 1998 to 2023"
    assert sheet["A5"].value=="Data published September 26, 2024"
    archive_headers.append({"sheet":sheet.title,"title":sheet["A1"].value,"unit":sheet["A2"].value,
      "availableYears":sheet["A3"].value,"releaseText":sheet["A5"].value,"latestYear":max(years),
      "latestYearCell":f'{openpyxl.utils.get_column_letter(sheet.max_column)}8',"contains2025":False})

latest_headers=[]
workbook=openpyxl.load_workbook(HERE.parent/"us/ValueAdded-2026-09-14.xlsx",data_only=True)
for name in ["TVA112-A","TVA113-A"]:
    sheet=workbook[name]
    years=[int(c.value) for c in sheet[8] if str(c.value).isdigit()]
    assert max(years)==2024 and 2025 not in years
    latest_headers.append({"sourceFile":"../us/ValueAdded-2026-09-14.xlsx","sheet":name,
      "title":sheet["A1"].value,"periodText":sheet["A3"].value,"releaseText":sheet["A5"].value,
      "latestYear":2024,"contains2025":False,"locator":"A1,A3,A5,D8:AE8"})

def body_text(file):
    source=(HERE/file).read_text()
    return re.sub(r"\s+"," ",unescape(re.sub(r"<[^>]+>"," ",source)))

discontinued=body_text("discontinued-statistics.html")
phrase="Components of Value Added"
pos=discontinued.index(phrase)
stop_excerpt=discontinued[pos-70:pos+len(phrase)+80]
assert "9/25/2025" in stop_excerpt and "Discontinued" in stop_excerpt
archive_page=body_text("industry-tables-no-longer-published.html")
assert "Last published September 26, 2024" in archive_page
update=body_text("annual-update.html")
assert "On September 30, 2026" in update
assert "All Industry Economic Accounts tables, including underlying detail" in update
release=body_text("gdp-release-status.html")
assert "Next release: September 30, 2026" in release

info=openpyxl.load_workbook(HERE/"ComponentsOfVaInfo.xlsx",data_only=True)["Definitions"]
formula={"expression":"VA = compensation of employees + taxes on production and imports less subsidies + gross operating surplus",
   "validCondition":"same year, industry, geography, price basis, and vintage for all three components; use GOS including consumption of fixed capital, not net surplus",
   "sourceFile":"ComponentsOfVaInfo.xlsx","locator":"Definitions!A4:B15 (hierarchy); A27,A29,A35,A37 (definitions)",
   "2025Computable":False,"reason":"No 2025 component observations in continuing table TVA113 or discontinued detailed archive; cannot replace absent GOS or taxes with old shares or profits."}
assert "GOS" in info["A13"].value

api=json.loads((HERE/"api-empty-userid-response.txt").read_text())
assert api["BEAAPI"]["Results"]["Error"]["APIErrorCode"]=="1"
api_state={"status":"authentication-unavailable","actualResponse":api["BEAAPI"]["Results"]["Error"],
    "file":"api-empty-userid-response.txt","interpretation":"API Year metadata not verified with an authenticated API request; this error proves no data-availability claim. Official iTable data and year lists independently provide the availability evidence."}

candidate=json.loads((HERE.parent/"us/us-annual-va-candidate.json").read_text())
missing=[o for o in candidate["observations"] if o["year"]==2025 and o["value"] is None]
assert len(missing)==80
nodes={n["id"]:n for n in candidate["nodes"]}
row_meta=json.loads(prompts("itable-uva-current")["Rows"]["PromtData"])["Table"]
by_line={int(r["RowSeq"])-2:r for r in row_meta if str(r["RowSeq"]).isdigit()}
gaps=[]
for obs in missing:
    node=nodes[obs["nodeId"]]
    official=by_line[node["beaUnderlyingLine"]]
    assert node["nameEn"] in official["StubText"]
    gaps.append({"nodeId":node["id"],"name":node["name"],"nameEn":node["nameEn"],
      "parentIndustryId":node["parentIndustryId"],"officialIndustryCode":official["Code"],"year":2025,
      "directValueAdded":None,"compensation":None,"netProductionTaxes":None,"grossOperatingSurplus":None,
      "status":"official-target-year-not-published","checkedAt":CHECKED,
      "evidence":[{"file":"itable-uva-current-response.json","locator":"First_Year/Last_Year option maximum 2024; TheTable year header max 2024"},
                  {"file":"itable-components-current-response.json","locator":"TVA113 year header and option maximum 2024"},
                  {"file":"ComponentsOfVA.xlsx","locator":"UVCT2-A, UVCT3-A, UVCT40-A: A3/A5 and D8:AC8, latest 2023"}]})

report={"checkedAt":CHECKED,"scope":"80 missing US 2025 detail nodes only; no site edits or re-estimation",
    "underlyingDirectVA":direct,"currentComponents":components,"archivedComponentsHeaders":archive_headers,
    "continuingWorkbookHeaders":latest_headers,
    "forced2025Probe":{"requestFile":"itable-2025-request.json","responseFile":"itable-2025-response.json",
       "valid2025Table":False,"malformedTableJSON":malformed,"maximumYearStill2024":True,
       "warning":"The out-of-range response contains residual 2024 values but lacks a valid 2025 year header. Reject it; requested year is not proof of observed year."},
    "publicAPI":api_state,
    "archivedComponentsDiscontinued":{"announced":"2025-09-25","lastPublished":"2024-09-26","latestDataYear":2023,
       "source":"https://www.bea.gov/data/discontinued-or-delayed-statistics","locator":"National / Industry / Underlying detail Components of Value Added row",
       "excerpt":stop_excerpt},
    "componentsMethod":formula,
    "nextScheduledIndustryUpdate":{"date":"2026-09-30","source":"https://www.bea.gov/information-updates-national-regional-economic-accounts",
       "locator":"Data availability / On September 30, 2026 / All Industry Economic Accounts tables, including underlying detail",
       "isAlreadyReleased":False,"isGuaranteeOfEveryRequested2025Node":False},
    "unblockConditions":["BEA releases 2025 observations in UVA205 / UnderlyingGDPbyIndustry with matching industry codes, current-dollar units, year headers and revision metadata.",
       "Alternatively, BEA publishes all three 2025 VA components at those exact industry codes, with compatible scope and vintage; then compute VA as an explicitly labelled calculation.",
       "The scheduled date arriving, an API key becoming available, or a network retry alone is not evidence that the missing 2025 data have been published.",
       "The discontinued component archive requires an explicit resumption/new official series; the ordinary September 30 schedule does not imply its revival."],
    "missingNodes":gaps}
(HERE/"availability-audit.json").write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n")
manifest=[{"file":f.name,"bytes":f.stat().st_size,"sha256":hashlib.sha256(f.read_bytes()).hexdigest()} for f in sorted(HERE.iterdir()) if f.is_file() and f.name!="files-manifest.json"]
(HERE/"files-manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n")
print(json.dumps({"missing2025Nodes":len(gaps),"iTableLatestYear":2024,"currentComponentsLatestYear":2024,
   "archivedComponentsLatestYear":2023,"archiveDiscontinued":True,"nextScheduledIndustryUpdate":"2026-09-30",
   "publicAPICheck":"authentication unavailable; no availability inference","new2025Values":0}))
